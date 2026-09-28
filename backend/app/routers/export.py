"""
Endpoints d'export : articles (CSV/xlsx) et rapports de synthèse (PDF).
Toutes les routes exigent une connexion (get_current_user).

Choix PDF : ReportLab (pur Python) — pas de dépendance système (contrairement à
WeasyPrint qui nécessite GTK/Cairo, pénible sous Windows). Les polices intégrées
(Helvetica) gèrent nativement les accents français (encodage WinAnsi/Latin-1),
donc aucun enregistrement de police TTF n'est nécessaire.
"""

import csv
import io
import logging
from datetime import date, datetime, time, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

import app.models as models
from app.models import CompetitorEnum, AlertCategoryEnum, SentimentEnum
from app.database import get_db
from app.auth.dependencies import get_current_user

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/export", tags=["Export"])

EXPORT_COLUMNS = ["Titre", "Concurrent", "Catégorie", "Sentiment", "Résumé",
                  "Date de publication", "Source", "URL"]

_SENTIMENT_POLARITY = {"Positif": 1, "Neutre": 0, "Négatif": -1}


# ─── Filtrage partagé ─────────────────────────────────────────────────

def _filtered_query(db, competitor=None, category=None, sentiment=None, date_from=None, date_to=None):
    q = db.query(models.Article)
    if competitor is not None:
        q = q.filter(models.Article.competitor == competitor)
    if category is not None:
        q = q.filter(models.Article.alert_category == category)
    if sentiment is not None:
        q = q.filter(models.Article.sentiment == sentiment)
    if date_from is not None:
        q = q.filter(models.Article.published_at >= datetime.combine(date_from, time.min))
    if date_to is not None:
        q = q.filter(models.Article.published_at <= datetime.combine(date_to, time.max))
    return q.order_by(models.Article.published_at.desc())


def _article_row(a: models.Article):
    return [
        a.title or "",
        a.competitor.value if a.competitor else "",
        a.alert_category.value if a.alert_category else "Non classé",
        a.sentiment.value if a.sentiment else "",
        a.summary or "",
        a.published_at.strftime("%Y-%m-%d %H:%M") if a.published_at else "",
        a.source or "",
        a.url or "",
    ]


def _sentiment_label(articles):
    vals = [_SENTIMENT_POLARITY[a.sentiment.value] for a in articles if a.sentiment]
    if not vals:
        return "—"
    avg = sum(vals) / len(vals)
    tone = "Positif" if avg > 0.1 else "Négatif" if avg < -0.1 else "Neutre"
    return f"{tone} ({avg:+.2f})"


# ─── Export articles (CSV / xlsx) ─────────────────────────────────────

@router.get("/articles")
def export_articles(
    format: str = Query("csv", pattern="^(csv|xlsx)$"),
    competitor: Optional[CompetitorEnum] = None,
    category: Optional[AlertCategoryEnum] = None,
    sentiment: Optional[SentimentEnum] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    articles = _filtered_query(db, competitor, category, sentiment, date_from, date_to).all()
    today = date.today().isoformat()

    if format == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";")  # ; = séparateur attendu par Excel FR
        writer.writerow(EXPORT_COLUMNS)
        for a in articles:
            writer.writerow(_article_row(a))
        # utf-8-sig ajoute le BOM pour qu'Excel affiche correctement les accents.
        content = buffer.getvalue().encode("utf-8-sig")
        return StreamingResponse(
            io.BytesIO(content),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="articles_{today}.csv"'},
        )

    # xlsx via openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Articles"
    ws.append(EXPORT_COLUMNS)
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for a in articles:
        ws.append(_article_row(a))
    # Largeurs de colonnes lisibles
    for col, width in zip("ABCDEFGH", (40, 16, 18, 12, 60, 18, 20, 40)):
        ws.column_dimensions[col].width = width

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="articles_{today}.xlsx"'},
    )


# ─── Génération PDF (ReportLab) ───────────────────────────────────────

def _pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("ReportTitle", parent=styles["Title"], fontSize=20,
                              textColor=colors.HexColor("#0f172a"), spaceAfter=4))
    styles.add(ParagraphStyle("SubTitle", parent=styles["Normal"], fontSize=9,
                              textColor=colors.HexColor("#64748b"), spaceAfter=14))
    styles.add(ParagraphStyle("Section", parent=styles["Heading2"], fontSize=13,
                              textColor=colors.HexColor("#0284c7"), spaceBefore=14, spaceAfter=6))
    styles.add(ParagraphStyle("Cell", parent=styles["Normal"], fontSize=8, leading=10))
    styles.add(ParagraphStyle("CellHead", parent=styles["Normal"], fontSize=8, leading=10,
                              textColor=colors.white, fontName="Helvetica-Bold"))
    styles.add(ParagraphStyle("Note", parent=styles["Normal"], fontSize=10,
                              textColor=colors.HexColor("#64748b"), spaceBefore=6))
    return styles


def _table(rows, col_widths, styles):
    """rows : liste de listes de chaînes ; 1ʳᵉ ligne = en-tête colorée."""
    data = []
    for i, row in enumerate(rows):
        style = styles["CellHead"] if i == 0 else styles["Cell"]
        data.append([Paragraph(str(c), style) for c in row])
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    return t


def _build_pdf(story, filename):
    bio = io.BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            leftMargin=1.6 * cm, rightMargin=1.6 * cm, title=filename)
    doc.build(story)
    bio.seek(0)
    return StreamingResponse(
        bio, media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _competitor_section(story, styles, competitor_value, articles, heading_level="Section", full=True):
    """Ajoute au story les stats d'un concurrent. full=True inclut critiques + 10 derniers."""
    if not articles:
        story.append(Paragraph("Aucune donnée sur la période.", styles["Note"]))
        return

    total = len(articles)
    # KPIs
    kpi = [["Indicateur", "Valeur"],
           ["Nombre d'articles", str(total)],
           ["Sentiment moyen", _sentiment_label(articles)]]
    story.append(_table(kpi, [8 * cm, 8 * cm], styles))
    story.append(Spacer(1, 8))

    # Répartition par catégorie
    cat_counts = {}
    for a in articles:
        k = a.alert_category.value if a.alert_category else "Non classé"
        cat_counts[k] = cat_counts.get(k, 0) + 1
    cat_rows = [["Catégorie", "Nombre"]] + [[k, str(v)] for k, v in sorted(cat_counts.items(), key=lambda x: -x[1])]
    story.append(Paragraph("Répartition par catégorie", styles["Section"]))
    story.append(_table(cat_rows, [10 * cm, 6 * cm], styles))

    if not full:
        return

    # Alertes critiques
    story.append(Paragraph("Alertes critiques", styles["Section"]))
    critical = [a for a in articles if a.is_critical]
    if critical:
        crit_rows = [["Date", "Catégorie", "Titre"]] + [[
            a.published_at.strftime("%Y-%m-%d") if a.published_at else "—",
            a.alert_category.value if a.alert_category else "Non classé",
            a.title or "",
        ] for a in critical]
        story.append(_table(crit_rows, [2.5 * cm, 3.5 * cm, 10 * cm], styles))
    else:
        story.append(Paragraph("Aucune alerte critique.", styles["Note"]))

    # 10 derniers articles
    story.append(Paragraph("10 derniers articles", styles["Section"]))
    last_rows = [["Date", "Catégorie", "Titre & résumé"]]
    for a in articles[:10]:
        resume = (a.summary or "").strip()
        if len(resume) > 220:
            resume = resume[:220] + "…"
        titre_resume = f"<b>{(a.title or '')}</b><br/>{resume}" if resume else f"<b>{(a.title or '')}</b>"
        last_rows.append([
            a.published_at.strftime("%Y-%m-%d") if a.published_at else "—",
            a.alert_category.value if a.alert_category else "Non classé",
            titre_resume,
        ])
    story.append(_table(last_rows, [2.5 * cm, 3.5 * cm, 10 * cm], styles))


@router.get("/report/{competitor}")
def export_competitor_report(
    competitor: str,
    format: str = Query("pdf", pattern="^pdf$"),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    # Concurrent inconnu -> 404 (plutôt que 422)
    try:
        comp = CompetitorEnum(competitor)
    except ValueError:
        raise HTTPException(status_code=404, detail=f"Concurrent inconnu : {competitor}")

    articles = _filtered_query(db, competitor=comp).all()
    styles = _pdf_styles()
    story = [
        Paragraph(f"Rapport de veille — {comp.value}", styles["ReportTitle"]),
        Paragraph(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}", styles["SubTitle"]),
    ]
    _competitor_section(story, styles, comp.value, articles, full=True)

    safe = comp.value.replace(" ", "_")
    return _build_pdf(story, f"rapport_{safe}_{date.today().isoformat()}.pdf")


@router.get("/weekly-report")
def export_weekly_report(
    format: str = Query("pdf", pattern="^pdf$"),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    since = datetime.utcnow() - timedelta(days=7)
    styles = _pdf_styles()
    story = [
        Paragraph("Rapport hebdomadaire de veille", styles["ReportTitle"]),
        Paragraph(
            f"7 derniers jours — généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}",
            styles["SubTitle"],
        ),
    ]

    all_week = db.query(models.Article).filter(models.Article.published_at >= since).all()
    if not all_week:
        story.append(Paragraph("Aucune donnée sur les 7 derniers jours.", styles["Note"]))
        return _build_pdf(story, f"rapport_hebdo_{date.today().isoformat()}.pdf")

    for comp in CompetitorEnum:
        arts = [a for a in all_week if a.competitor == comp]
        story.append(Paragraph(comp.value, styles["Section"]))
        # section allégée par concurrent (KPIs + catégories), sans les 10 derniers
        _competitor_section(story, styles, comp.value, arts, full=False)

    return _build_pdf(story, f"rapport_hebdo_{date.today().isoformat()}.pdf")
